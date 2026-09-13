from __future__ import annotations

import importlib
import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Any


class DocHistoryUnavailable(RuntimeError):
    """Raised when the reusable doc_history package cannot be imported."""


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _candidate_doc_history_roots() -> tuple[Path, ...]:
    """Return local locations where the sibling doc_history repo may live.

    DiScO does not copy doc_history's extraction logic. It imports the package
    directly, preferring a normal installed import and then looking for the
    user's local sibling checkout.
    """

    candidates: list[Path] = []
    configured = os.environ.get("DOC_HISTORY_PATH")
    if configured:
        candidates.append(Path(configured).expanduser())

    repo_root = _repo_root()
    candidates.extend(
        [
            repo_root.parent / "doc_history",
            Path.home() / "doc_history",
        ]
    )

    # Preserve order while removing duplicates.
    unique: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate.resolve()) if candidate.exists() else str(candidate)
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    return tuple(unique)


def _load_doc_history_module():
    try:
        return importlib.import_module("doc_history")
    except ModuleNotFoundError:
        pass

    for root in _candidate_doc_history_roots():
        if not (root / "doc_history" / "__init__.py").exists():
            continue
        root_text = str(root.resolve())
        if root_text not in sys.path:
            sys.path.insert(0, root_text)
        try:
            return importlib.import_module("doc_history")
        except ModuleNotFoundError:
            continue

    searched = ", ".join(str(path) for path in _candidate_doc_history_roots())
    raise DocHistoryUnavailable(
        "DiScO could not import doc_history. Clone/install natosit-dev/doc_history "
        f"or set DOC_HISTORY_PATH. Searched: {searched}"
    )


def inspect_document_artifact(data: bytes, filename: str) -> dict[str, Any]:
    """Reuse doc_history and return its complete artifact/provenance dataset."""

    module = _load_doc_history_module()
    result = module.inspect_bytes(data, filename)
    return {
        "source_type": "file",
        "filename": filename,
        "doc_history": result,
        "timeline_events": module.timeline_events(result),
        "provenance_clues": module.provenance_clues(result),
    }


def extract_document_text(data: bytes, filename: str) -> str:
    """Extract text for DiScO scoring without changing doc_history metadata.

    doc_history remains the provenance/metadata authority. This helper only
    supplies the plain text consumed by the existing DiScO text engine.
    """

    suffix = Path(filename).suffix.lower()

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(data))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text:
                blocks.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append("\t".join(cells))
        return "\n".join(blocks)

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data), strict=False)
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    raise ValueError("DiScO file upload currently supports .docx and .pdf files.")
